"""Cover letter generation service."""
import io
import logging

from docx import Document
from docx.shared import Pt, Inches

from backend.services.ai_service import ai_service, AIServiceError
from backend.utils.templates import COVER_LETTER_SYSTEM_PROMPT

logger = logging.getLogger(__name__)


def generate_cover_letter(profile_data: dict, job_description: str, parsed_job: dict, tone: str = "formal") -> str:
    """Generate a tailored cover letter using AI."""
    profile_text = _format_profile_for_cover_letter(profile_data)

    try:
        result = ai_service.generate(
            system_prompt=COVER_LETTER_SYSTEM_PROMPT,
            user_prompt=(
                f"Candidate Profile:\n{profile_text}\n\n"
                f"Job Description:\n{job_description}\n\n"
                f"Parsed Job Info:\n"
                f"Job Title: {parsed_job.get('job_title', 'N/A')}\n"
                f"Company: {parsed_job.get('company_name', 'N/A')}\n"
                f"Required Skills: {', '.join(parsed_job.get('required_skills', []))}\n"
                f"Key Responsibilities: {', '.join(parsed_job.get('key_responsibilities', []))}\n\n"
                f"Tone: {tone}\n\n"
                f"Generate a professional cover letter."
            ),
            response_format="json",
        )
        return result.get("cover_letter", result.get("content", ""))
    except AIServiceError as e:
        logger.error(f"Cover letter generation failed: {e}")
        raise ValueError("Failed to generate cover letter. The AI service is temporarily unavailable. Please try again.")


def generate_cover_letter_docx(content: str, user_info: dict) -> io.BytesIO:
    """Generate a DOCX file from cover letter content."""
    doc = Document()

    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Header with contact info
    name = user_info.get("name", "")
    if name:
        name_para = doc.add_paragraph()
        name_run = name_para.add_run(name)
        name_run.bold = True
        name_run.font.size = Pt(14)

    contact_parts = []
    for field in ["email", "phone", "location"]:
        if user_info.get(field):
            contact_parts.append(user_info[field])
    if contact_parts:
        contact_para = doc.add_paragraph(" | ".join(contact_parts))
        contact_para.space_after = Pt(12)

    # Cover letter body - split by double newlines into paragraphs
    paragraphs = content.strip().split("\n\n")
    for para_text in paragraphs:
        clean_text = para_text.strip()
        if clean_text:
            doc.add_paragraph(clean_text)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def _format_profile_for_cover_letter(profile: dict) -> str:
    """Format profile data for cover letter generation."""
    lines = []
    user = profile.get("user", {})
    lines.append(f"Name: {user.get('name', 'N/A')}")
    lines.append(f"Summary: {user.get('professional_summary', 'N/A')}")

    lines.append("\nKey Experience:")
    for exp in profile.get("experiences", [])[:3]:
        lines.append(f"  - {exp['job_title']} at {exp['company_name']}")
        for bullet in (exp.get("bullet_points") or [])[:2]:
            lines.append(f"    * {bullet}")

    skills = [s["skill_name"] for s in profile.get("skills", [])]
    if skills:
        lines.append(f"\nSkills: {', '.join(skills)}")

    for edu in profile.get("education", [])[:1]:
        lines.append(f"\nEducation: {edu['degree']} from {edu['institution']}")

    return "\n".join(lines)

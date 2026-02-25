"""Resume generation service - builds tailored resumes using AI and exports to DOCX/HTML."""
import io
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from jinja2 import Template

from backend.services.ai_service import ai_service
from backend.utils.templates import (
    RESUME_GENERATOR_SYSTEM_PROMPT, RESUME_HTML_TEMPLATE,
    SECTION_REGENERATE_PROMPT, HTML_TEMPLATES,
)


def generate_tailored_resume(profile_data: dict, job_description: str, parsed_job: dict) -> dict:
    """Generate a tailored resume using AI, based on profile and job description."""
    profile_text = _format_profile_for_resume(profile_data)

    result = ai_service.generate(
        system_prompt=RESUME_GENERATOR_SYSTEM_PROMPT,
        user_prompt=(
            f"Candidate's Complete Profile:\n{profile_text}\n\n"
            f"Job Description:\n{job_description}\n\n"
            f"Parsed Job Data:\n"
            f"Required Skills: {', '.join(parsed_job.get('required_skills', []))}\n"
            f"Preferred Skills: {', '.join(parsed_job.get('preferred_skills', []))}\n"
            f"ATS Keywords: {', '.join(parsed_job.get('ats_keywords', []))}\n\n"
            f"Generate a tailored ATS-optimized resume using ONLY the candidate's real experience."
        ),
        response_format="json",
    )
    return result


def generate_resume_html(resume_data: dict, user_info: dict, template_name: str = "ats_classic") -> str:
    """Render the generated resume as HTML for preview."""
    contact_parts = []
    if user_info.get("email"):
        contact_parts.append(user_info["email"])
    if user_info.get("phone"):
        contact_parts.append(user_info["phone"])
    if user_info.get("location"):
        contact_parts.append(user_info["location"])
    if user_info.get("linkedin_url"):
        contact_parts.append(f'<a href="{user_info["linkedin_url"]}">LinkedIn</a>')
    if user_info.get("portfolio_url"):
        contact_parts.append(f'<a href="{user_info["portfolio_url"]}">Portfolio</a>')

    html_template = HTML_TEMPLATES.get(template_name, RESUME_HTML_TEMPLATE)
    template = Template(html_template)
    html = template.render(
        name=user_info.get("name", ""),
        contact_line=" | ".join(contact_parts),
        professional_summary=resume_data.get("professional_summary", ""),
        work_experience=resume_data.get("work_experience", []),
        education=resume_data.get("education", []),
        skills=resume_data.get("skills", {}),
        projects=resume_data.get("projects", []),
        certifications=resume_data.get("certifications", []),
    )
    return html


def generate_resume_docx(resume_data: dict, user_info: dict) -> io.BytesIO:
    """Generate an ATS-friendly .docx file from resume data."""
    doc = Document()

    # Set default font
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)

    # Set narrow margins
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Name header
    name_para = doc.add_paragraph()
    name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = name_para.add_run(user_info.get("name", ""))
    name_run.bold = True
    name_run.font.size = Pt(18)

    # Contact info
    contact_parts = []
    for field in ["email", "phone", "location", "linkedin_url", "portfolio_url"]:
        if user_info.get(field):
            contact_parts.append(user_info[field])
    if contact_parts:
        contact_para = doc.add_paragraph()
        contact_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        contact_run = contact_para.add_run(" | ".join(contact_parts))
        contact_run.font.size = Pt(10)

    # Professional Summary
    if resume_data.get("professional_summary"):
        _add_section_header(doc, "Professional Summary")
        doc.add_paragraph(resume_data["professional_summary"])

    # Work Experience
    if resume_data.get("work_experience"):
        _add_section_header(doc, "Work Experience")
        for job in resume_data["work_experience"]:
            job_para = doc.add_paragraph()
            title_run = job_para.add_run(f"{job.get('title', '')} | {job.get('company', '')}")
            title_run.bold = True
            if job.get("location"):
                job_para.add_run(f" | {job['location']}")
            if job.get("dates"):
                job_para.add_run(f"\t{job['dates']}")

            for bullet in job.get("bullets", []):
                bullet_para = doc.add_paragraph(style="List Bullet")
                bullet_para.text = bullet

    # Education
    if resume_data.get("education"):
        _add_section_header(doc, "Education")
        for edu in resume_data["education"]:
            edu_para = doc.add_paragraph()
            edu_run = edu_para.add_run(f"{edu.get('degree', '')}")
            edu_run.bold = True
            edu_para.add_run(f" | {edu.get('institution', '')}")
            if edu.get("graduation_date"):
                edu_para.add_run(f" | {edu['graduation_date']}")
            if edu.get("gpa"):
                doc.add_paragraph(f"GPA: {edu['gpa']}")

    # Technical Skills
    if resume_data.get("skills"):
        _add_section_header(doc, "Technical Skills")
        for category, skill_list in resume_data["skills"].items():
            if skill_list:
                skill_para = doc.add_paragraph()
                cat_run = skill_para.add_run(f"{category}: ")
                cat_run.bold = True
                skill_para.add_run(", ".join(skill_list))

    # Projects
    if resume_data.get("projects"):
        _add_section_header(doc, "Projects")
        for project in resume_data["projects"]:
            proj_para = doc.add_paragraph()
            proj_run = proj_para.add_run(project.get("name", ""))
            proj_run.bold = True
            if project.get("technologies"):
                proj_para.add_run(f" | {', '.join(project['technologies'])}")
            if project.get("url"):
                proj_para.add_run(f" | {project['url']}")

            for bullet in project.get("bullets", []):
                bullet_para = doc.add_paragraph(style="List Bullet")
                bullet_para.text = bullet

    # Certifications
    if resume_data.get("certifications"):
        _add_section_header(doc, "Certifications")
        for cert in resume_data["certifications"]:
            cert_para = doc.add_paragraph()
            cert_run = cert_para.add_run(cert.get("name", ""))
            cert_run.bold = True
            cert_para.add_run(f" — {cert.get('issuing_org', '')}")
            if cert.get("date"):
                cert_para.add_run(f" ({cert['date']})")

    # Save to BytesIO
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def _add_section_header(doc: Document, text: str):
    """Add an ATS-friendly section header with underline."""
    para = doc.add_paragraph()
    para.space_before = Pt(8)
    para.space_after = Pt(4)
    run = para.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(12)
    # Add a bottom border effect
    from docx.oxml.ns import qn
    pPr = para._element.get_or_add_pPr()
    pBdr = pPr.makeelement(qn("w:pBdr"), {})
    bottom = pBdr.makeelement(
        qn("w:bottom"),
        {qn("w:val"): "single", qn("w:sz"): "6", qn("w:space"): "1", qn("w:color"): "333333"},
    )
    pBdr.append(bottom)
    pPr.append(pBdr)


def generate_resume_docx_with_template(resume_data: dict, user_info: dict, template_name: str = "ats_classic") -> io.BytesIO:
    """Generate DOCX with template selection. Modern and compact have different styling."""
    if template_name == "modern":
        return _generate_modern_docx(resume_data, user_info)
    elif template_name == "compact":
        return _generate_compact_docx(resume_data, user_info)
    return generate_resume_docx(resume_data, user_info)


def _generate_modern_docx(resume_data: dict, user_info: dict) -> io.BytesIO:
    """Modern template with blue accent color."""
    from docx.shared import RGBColor
    doc = Document()

    style = doc.styles["Normal"]
    font = style.font
    font.name = "Segoe UI"
    font.size = Pt(11)

    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Name in blue
    name_para = doc.add_paragraph()
    name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = name_para.add_run(user_info.get("name", ""))
    name_run.bold = True
    name_run.font.size = Pt(20)
    name_run.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)

    # Contact info
    contact_parts = [user_info.get(f) for f in ["email", "phone", "location", "linkedin_url", "portfolio_url"] if user_info.get(f)]
    if contact_parts:
        contact_para = doc.add_paragraph()
        contact_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        contact_run = contact_para.add_run(" | ".join(contact_parts))
        contact_run.font.size = Pt(10)
        contact_run.font.color.rgb = RGBColor(0x71, 0x80, 0x96)

    def add_modern_header(text):
        para = doc.add_paragraph()
        para.space_before = Pt(10)
        para.space_after = Pt(4)
        run = para.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)

    # Sections follow same pattern as classic but with blue headers
    if resume_data.get("professional_summary"):
        add_modern_header("Professional Summary")
        doc.add_paragraph(resume_data["professional_summary"])

    if resume_data.get("work_experience"):
        add_modern_header("Work Experience")
        for job in resume_data["work_experience"]:
            job_para = doc.add_paragraph()
            title_run = job_para.add_run(f"{job.get('title', '')} | {job.get('company', '')}")
            title_run.bold = True
            if job.get("dates"):
                job_para.add_run(f"\t{job['dates']}")
            for bullet in job.get("bullets", []):
                bp = doc.add_paragraph(style="List Bullet")
                bp.text = bullet

    if resume_data.get("education"):
        add_modern_header("Education")
        for edu in resume_data["education"]:
            edu_para = doc.add_paragraph()
            edu_run = edu_para.add_run(edu.get("degree", ""))
            edu_run.bold = True
            edu_para.add_run(f" | {edu.get('institution', '')}")
            if edu.get("graduation_date"):
                edu_para.add_run(f" | {edu['graduation_date']}")

    if resume_data.get("skills"):
        add_modern_header("Technical Skills")
        for category, skill_list in resume_data["skills"].items():
            if skill_list:
                sp = doc.add_paragraph()
                cr = sp.add_run(f"{category}: ")
                cr.bold = True
                sp.add_run(", ".join(skill_list))

    if resume_data.get("projects"):
        add_modern_header("Projects")
        for project in resume_data["projects"]:
            pp = doc.add_paragraph()
            pr = pp.add_run(project.get("name", ""))
            pr.bold = True
            if project.get("technologies"):
                pp.add_run(f" | {', '.join(project['technologies'])}")
            for bullet in project.get("bullets", []):
                bp = doc.add_paragraph(style="List Bullet")
                bp.text = bullet

    if resume_data.get("certifications"):
        add_modern_header("Certifications")
        for cert in resume_data["certifications"]:
            cp = doc.add_paragraph()
            cr = cp.add_run(cert.get("name", ""))
            cr.bold = True
            cp.add_run(f" — {cert.get('issuing_org', '')}")
            if cert.get("date"):
                cp.add_run(f" ({cert['date']})")

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def _generate_compact_docx(resume_data: dict, user_info: dict) -> io.BytesIO:
    """Compact template with tighter spacing."""
    doc = Document()

    style = doc.styles["Normal"]
    font = style.font
    font.name = "Arial"
    font.size = Pt(10)

    for section in doc.sections:
        section.top_margin = Inches(0.4)
        section.bottom_margin = Inches(0.4)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    # Name
    name_para = doc.add_paragraph()
    name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_para.space_after = Pt(2)
    name_run = name_para.add_run(user_info.get("name", ""))
    name_run.bold = True
    name_run.font.size = Pt(15)

    # Contact
    contact_parts = [user_info.get(f) for f in ["email", "phone", "location", "linkedin_url", "portfolio_url"] if user_info.get(f)]
    if contact_parts:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.space_after = Pt(4)
        cr = cp.add_run(" | ".join(contact_parts))
        cr.font.size = Pt(9)

    def add_compact_header(text):
        para = doc.add_paragraph()
        para.space_before = Pt(6)
        para.space_after = Pt(2)
        run = para.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(10)
        # thin bottom border
        from docx.oxml.ns import qn
        pPr = para._element.get_or_add_pPr()
        pBdr = pPr.makeelement(qn("w:pBdr"), {})
        bottom = pBdr.makeelement(qn("w:bottom"), {qn("w:val"): "single", qn("w:sz"): "4", qn("w:space"): "1", qn("w:color"): "999999"})
        pBdr.append(bottom)
        pPr.append(pBdr)

    if resume_data.get("professional_summary"):
        add_compact_header("Summary")
        p = doc.add_paragraph(resume_data["professional_summary"])
        p.paragraph_format.space_after = Pt(2)

    if resume_data.get("work_experience"):
        add_compact_header("Experience")
        for job in resume_data["work_experience"]:
            jp = doc.add_paragraph()
            jp.paragraph_format.space_after = Pt(1)
            tr = jp.add_run(f"{job.get('title', '')} | {job.get('company', '')}")
            tr.bold = True
            tr.font.size = Pt(10)
            if job.get("dates"):
                jp.add_run(f"\t{job['dates']}")
            for bullet in job.get("bullets", []):
                bp = doc.add_paragraph(style="List Bullet")
                bp.text = bullet
                bp.paragraph_format.space_after = Pt(0)

    if resume_data.get("education"):
        add_compact_header("Education")
        for edu in resume_data["education"]:
            ep = doc.add_paragraph()
            er = ep.add_run(edu.get("degree", ""))
            er.bold = True
            ep.add_run(f" | {edu.get('institution', '')} | {edu.get('graduation_date', '')}")

    if resume_data.get("skills"):
        add_compact_header("Skills")
        for category, skill_list in resume_data["skills"].items():
            if skill_list:
                sp = doc.add_paragraph()
                sp.paragraph_format.space_after = Pt(1)
                cr = sp.add_run(f"{category}: ")
                cr.bold = True
                sp.add_run(", ".join(skill_list))

    if resume_data.get("projects"):
        add_compact_header("Projects")
        for project in resume_data["projects"]:
            pp = doc.add_paragraph()
            pr = pp.add_run(project.get("name", ""))
            pr.bold = True
            if project.get("technologies"):
                pp.add_run(f" | {', '.join(project['technologies'])}")
            for bullet in project.get("bullets", []):
                bp = doc.add_paragraph(style="List Bullet")
                bp.text = bullet
                bp.paragraph_format.space_after = Pt(0)

    if resume_data.get("certifications"):
        add_compact_header("Certifications")
        for cert in resume_data["certifications"]:
            cp = doc.add_paragraph()
            cr = cp.add_run(cert.get("name", ""))
            cr.bold = True
            cp.add_run(f" — {cert.get('issuing_org', '')} ({cert.get('date', '')})")

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def regenerate_section(section_name: str, current_content: dict, job_description: str, profile_data: dict) -> dict:
    """Regenerate a single section of the resume using AI."""
    profile_text = _format_profile_for_resume(profile_data)
    prompt = SECTION_REGENERATE_PROMPT.format(
        section_name=section_name,
        current_content=str(current_content),
        job_description=job_description,
        profile_text=profile_text,
    )
    result = ai_service.generate(
        system_prompt="You are a professional resume writer. Return ONLY valid JSON.",
        user_prompt=prompt,
        response_format="json",
    )
    return result


def _format_profile_for_resume(profile: dict) -> str:
    """Format the full profile as text for the AI prompt."""
    lines = []
    user = profile.get("user", {})
    lines.append(f"Name: {user.get('name', 'N/A')}")
    lines.append(f"Location: {user.get('location', 'N/A')}")
    lines.append(f"Professional Summary: {user.get('professional_summary', 'N/A')}")

    lines.append("\n--- Work Experience ---")
    for exp in profile.get("experiences", []):
        end = exp.get("end_date") or "Present"
        lines.append(f"\n{exp['job_title']} at {exp['company_name']} ({exp.get('location', '')})")
        lines.append(f"  {exp.get('start_date', '')} - {end}")
        for bullet in exp.get("bullet_points", []):
            lines.append(f"  * {bullet}")
        if exp.get("skills_used"):
            lines.append(f"  Skills used: {', '.join(exp['skills_used'])}")

    lines.append("\n--- Education ---")
    for edu in profile.get("education", []):
        lines.append(f"\n{edu['degree']} - {edu['institution']} ({edu.get('graduation_date', 'N/A')})")
        if edu.get("gpa"):
            lines.append(f"  GPA: {edu['gpa']}")
        if edu.get("relevant_coursework"):
            lines.append(f"  Coursework: {', '.join(edu['relevant_coursework'])}")

    lines.append("\n--- Skills ---")
    for skill in profile.get("skills", []):
        lines.append(f"  {skill['skill_name']} - {skill.get('category', '')} ({skill.get('proficiency_level', '')})")

    lines.append("\n--- Projects ---")
    for proj in profile.get("projects", []):
        lines.append(f"\n{proj['project_name']}")
        if proj.get("description"):
            lines.append(f"  {proj['description']}")
        if proj.get("technologies_used"):
            lines.append(f"  Technologies: {', '.join(proj['technologies_used'])}")
        for bullet in proj.get("bullet_points", []):
            lines.append(f"  * {bullet}")

    lines.append("\n--- Certifications ---")
    for cert in profile.get("certifications", []):
        lines.append(f"  {cert['cert_name']} - {cert.get('issuing_org', '')} ({cert.get('date_obtained', 'N/A')})")

    return "\n".join(lines)

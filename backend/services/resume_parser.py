"""Resume import service - extracts text from PDF/DOCX and parses into structured profile data."""
import io
import logging

from backend.services.ai_service import ai_service, AIServiceError
from backend.utils.templates import RESUME_IMPORT_SYSTEM_PROMPT

logger = logging.getLogger(__name__)


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract all text from a DOCX file."""
    try:
        from docx import Document
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())
        return "\n".join(paragraphs)
    except Exception as e:
        logger.error(f"Failed to extract text from DOCX: {type(e).__name__}: {e}")
        raise ValueError("Could not read the DOCX file. It may be corrupted or password-protected.")


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract all text from a PDF file."""
    try:
        from PyPDF2 import PdfReader
        reader = PdfReader(io.BytesIO(file_bytes))
        pages = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                pages.append(text)
        if not pages:
            raise ValueError("No readable text found in PDF. It may be image-based or scanned.")
        return "\n".join(pages)
    except ValueError:
        raise
    except Exception as e:
        logger.error(f"Failed to extract text from PDF: {type(e).__name__}: {e}")
        raise ValueError("Could not read the PDF file. It may be corrupted or password-protected.")


def parse_resume_to_profile(text: str) -> dict:
    """Use AI to parse raw resume text into structured profile data."""
    try:
        result = ai_service.generate(
            system_prompt=RESUME_IMPORT_SYSTEM_PROMPT,
            user_prompt=f"Parse the following resume into structured profile data:\n\n{text}",
            response_format="json",
        )
        return result
    except AIServiceError as e:
        logger.error(f"AI resume parsing failed: {e}")
        raise ValueError("Failed to parse resume content. Please try again or enter your information manually.")
